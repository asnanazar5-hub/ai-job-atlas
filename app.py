"""AI Job Atlas - a Flask website for exploring AI jobs.

Run:  python app.py   then open http://127.0.0.1:5000
"""
import glob
import io
import os
import re

import pandas as pd
from docx import Document
from flask import Flask, jsonify, render_template, request
from pypdf import PdfReader

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # resumes up to 8 MB

# ----------------------------------------------------------------------------
# Lookup tables
# ----------------------------------------------------------------------------
LEVELS = {"EN": "Entry", "MI": "Mid", "SE": "Senior", "EX": "Executive"}
TYPES = {"FT": "Full-time", "PT": "Part-time", "CT": "Contract", "FL": "Freelance"}
MODES = {0: "On-site", 50: "Hybrid", 100: "Remote"}
EDU_RANK = {"Associate": 1, "Bachelor": 2, "Master": 3, "PhD": 4}
COUNTRY_CODES = {
    "Australia": "AU", "Austria": "AT", "Canada": "CA", "China": "CN",
    "Denmark": "DK", "Finland": "FI", "France": "FR", "Germany": "DE",
    "India": "IN", "Ireland": "IE", "Israel": "IL", "Japan": "JP",
    "Netherlands": "NL", "Norway": "NO", "Singapore": "SG",
    "South Korea": "KR", "Sweden": "SE", "Switzerland": "CH",
    "United Kingdom": "GB", "United States": "US",
}


def flag(country):
    """Turn a country name into its flag emoji."""
    code = COUNTRY_CODES.get(country, "")
    return "".join(chr(0x1F1E6 + ord(c) - 65) for c in code)


# ----------------------------------------------------------------------------
# Load and merge every CSV found in the data/ folder
# ----------------------------------------------------------------------------
def load_data():
    frames = []
    for n, path in enumerate(sorted(glob.glob(os.path.join(BASE_DIR, "data", "*.csv"))), 1):
        part = pd.read_csv(path)
        part["job_id"] = f"D{n}-" + part["job_id"].astype(str)  # ids repeat across files
        if "salary_local" not in part.columns:
            part["salary_local"] = part["salary_usd"]
        frames.append(part)
    df = pd.concat(frames, ignore_index=True)
    key = ["job_title", "salary_usd", "company_location", "posting_date", "company_name", "required_skills"]
    df = df.drop_duplicates(subset=key).reset_index(drop=True)

    df["skills_list"] = df["required_skills"].map(lambda s: [x.strip() for x in str(s).split(",") if x.strip()])
    df["skillset"] = df["skills_list"].map(frozenset)
    df["mode"] = df["remote_ratio"].map(MODES)
    df["search"] = (df["job_title"] + " " + df["company_name"] + " " + df["industry"] + " "
                    + df["company_location"] + " " + df["required_skills"]).str.lower()
    df["posted_ord"] = pd.to_datetime(df["posting_date"]).map(lambda t: t.toordinal())
    return df


DF = load_data()
ALL_SKILLS = sorted({s for js in DF["skillset"] for s in js})
ALL_TITLES = sorted(DF["job_title"].unique())
ALL_COUNTRIES = sorted(DF["company_location"].unique())
ALL_INDUSTRIES = sorted(DF["industry"].unique())


def build_country_cards():
    cards = []
    for country, g in DF.groupby("company_location"):
        cards.append({
            "name": country,
            "flag": flag(country),
            "code": COUNTRY_CODES.get(country, ""),
            "jobs": int(len(g)),
            "avg_salary": int(g["salary_usd"].mean()),
            "remote_pct": round(float((g["remote_ratio"] == 100).mean() * 100)),
            "top_title": g["job_title"].value_counts().index[0],
            "top_skill": g["skills_list"].explode().value_counts().index[0],
        })
    return cards


COUNTRY_CARDS = build_country_cards()


def meta():
    return {
        "countries": ALL_COUNTRIES, "titles": ALL_TITLES, "skills": ALL_SKILLS,
        "industries": ALL_INDUSTRIES, "levels": LEVELS, "types": TYPES,
        "flags": {c: flag(c) for c in ALL_COUNTRIES},
    }


# ----------------------------------------------------------------------------
# Pages
# ----------------------------------------------------------------------------
@app.get("/")
def home():
    top_roles = (DF.groupby("job_title")
                 .agg(jobs=("job_id", "count"), salary=("salary_usd", "mean"))
                 .sort_values("jobs", ascending=False).head(6).reset_index())
    hero = {
        "jobs": f"{len(DF):,}", "countries": len(ALL_COUNTRIES), "skills": len(ALL_SKILLS),
        "salary": f"${int(DF['salary_usd'].mean()):,}",
    }
    return render_template("home.html", hero=hero, skills=ALL_SKILLS,
                           roles=top_roles.to_dict("records"), page="home")


@app.get("/finder")
def finder():
    return render_template("finder.html", meta=meta(), page="finder")


@app.get("/countries")
def countries():
    return render_template("countries.html", cards=COUNTRY_CARDS, page="countries")


@app.get("/compare")
def compare():
    return render_template("compare.html", meta=meta(), page="compare")


@app.get("/insights")
def insights():
    return render_template("insights.html", meta=meta(), page="insights")


# ----------------------------------------------------------------------------
# Job search API: filters + match score + "your country first" ordering
# ----------------------------------------------------------------------------
STOP_WORDS = {"ai", "and", "of", "the", "a"}


def title_similarity(role, title):
    """1.0 for an exact title, otherwise word overlap between role and title."""
    t = title.lower()
    if role == t:
        return 1.0
    rw = {w for w in re.findall(r"[a-z]+", role) if w not in STOP_WORDS}
    tw = {w for w in re.findall(r"[a-z]+", t) if w not in STOP_WORDS}
    if not rw or not tw:
        return 0.0
    return len(rw & tw) / len(rw | tw)


@app.get("/api/jobs")
def api_jobs():
    a = request.args
    country = a.get("country", "").strip()

    # 1) filters --------------------------------------------------------------
    mask = pd.Series(True, index=DF.index)
    for col, key in (("experience_level", "level"), ("employment_type", "type"), ("industry", "industry")):
        if a.get(key):
            mask &= DF[col] == a[key]
    if a.get("title") in ALL_TITLES:
        mask &= DF["job_title"] == a["title"]
    if a.get("remote", "") in ("0", "50", "100"):
        mask &= DF["remote_ratio"] == int(a["remote"])
    q = a.get("q", "").strip().lower()
    if q:
        mask &= DF["search"].str.contains(q, regex=False)
    if a.get("scope") == "country" and country:
        mask &= DF["company_location"] == country
    sub = DF[mask].copy()

    # 2) match score (only when the user gave some profile information) --------
    skills = [s for s in a.get("skills", "").split(",") if s]
    uset = set(skills)
    role = a.get("role", "").strip().lower()
    parts = []  # (weight, scores)
    if uset:
        parts.append((0.6, sub["skillset"].map(lambda js: len(js & uset) / len(js))))
    if role:
        tmap = {t: title_similarity(role, t) for t in sub["job_title"].unique()}
        parts.append((0.2, sub["job_title"].map(tmap)))
    try:
        years = float(a["years"]) if a.get("years", "") != "" else None
    except ValueError:
        years = None
    if years is not None:
        parts.append((0.1, sub["years_experience"].map(
            lambda j: 1.0 if j <= years + 1 else max(0.0, 1 - (j - years - 1) / 5))))
    if a.get("education") in EDU_RANK:
        mine = EDU_RANK[a["education"]]
        parts.append((0.1, sub["education_required"].map(lambda e: 1.0 if EDU_RANK[e] <= mine else 0.3)))

    has_profile = bool(parts) and len(sub) > 0
    if has_profile:
        total_w = sum(w for w, _ in parts)
        sub["score"] = sum(w * s for w, s in parts) / total_w * 100
    else:
        sub["score"] = 0.0

    # 3) country first, then the chosen sort ---------------------------------
    # tier 2 = country you want to work in, tier 1 = the country you are from, tier 0 = everywhere else
    home = a.get("home", "").strip()
    sub["tier"] = 0
    if home:
        sub.loc[sub["company_location"] == home, "tier"] = 1
    if country:
        sub.loc[sub["company_location"] == country, "tier"] = 2
    sub["in_country"] = sub["tier"] == 2
    sort = a.get("sort", "match")
    if sort == "salary_desc":
        sub["k"] = sub["salary_usd"]
    elif sort == "salary_asc":
        sub["k"] = -sub["salary_usd"]
    elif sort == "recent":
        sub["k"] = sub["posted_ord"]
    else:
        sub["k"] = sub["score"] if has_profile else sub["posted_ord"]
    sub = sub.sort_values(["tier", "k", "salary_usd"], ascending=False)

    # 4) paginate -------------------------------------------------------------
    per = max(1, min(int(a.get("per_page", 20) or 20), 50))
    page = max(1, int(a.get("page", 1) or 1))
    total = len(sub)
    pages = max(1, -(-total // per))
    rows = sub.iloc[(page - 1) * per: page * per].to_dict("records")

    jobs = []
    for r in rows:
        jobs.append({
            "id": r["job_id"], "title": r["job_title"], "company": r["company_name"],
            "country": r["company_location"], "flag": flag(r["company_location"]),
            "residence": r["employee_residence"],
            "salary_usd": int(r["salary_usd"]), "salary_local": int(r["salary_local"]),
            "currency": r["salary_currency"],
            "level": LEVELS.get(r["experience_level"], r["experience_level"]),
            "type": TYPES.get(r["employment_type"], r["employment_type"]),
            "mode": r["mode"], "skills": r["skills_list"],
            "matched": [s for s in r["skills_list"] if s in uset],
            "education": r["education_required"], "years": int(r["years_experience"]),
            "industry": r["industry"], "size": r["company_size"],
            "posted": r["posting_date"], "deadline": r["application_deadline"],
            "benefits": float(r["benefits_score"]),
            "score": int(round(r["score"])) if has_profile else None,
            "in_country": bool(r["in_country"]), "tier": int(r["tier"]),
        })
    return jsonify({
        "total": total, "page": page, "pages": pages, "jobs": jobs,
        "in_country_total": int((sub["tier"] == 2).sum()),
        "home_total": int((sub["tier"] == 1).sum()),
        "other_total": int((sub["tier"] == 0).sum()),
        "has_profile": has_profile,
    })


# ----------------------------------------------------------------------------
# Insights API
# ----------------------------------------------------------------------------
@app.get("/api/stats")
def api_stats():
    country = request.args.get("country", "").strip()
    d = DF[DF["company_location"] == country] if country else DF
    if d.empty:
        return jsonify({"error": "No jobs for this country"}), 404

    titles = (d.groupby("job_title").agg(jobs=("job_id", "count"), salary=("salary_usd", "mean"))
              .sort_values("jobs", ascending=False).head(8).reset_index())
    levels = d.groupby("experience_level")["salary_usd"].mean()
    skills = d["skills_list"].explode().value_counts().head(12)
    industries = d["industry"].value_counts().head(8)
    modes = d["mode"].value_counts()
    edu = d.groupby("education_required")["salary_usd"].mean()
    months = d.assign(m=d["posting_date"].str[:7]).groupby("m").size().sort_index()

    out = {
        "country": country or "All countries",
        "kpi": {
            "jobs": int(len(d)), "avg_salary": int(d["salary_usd"].mean()),
            "median_salary": int(d["salary_usd"].median()),
            "remote_pct": round(float((d["remote_ratio"] == 100).mean() * 100)),
            "companies": int(d["company_name"].nunique()),
        },
        "titles": [{"label": r.job_title, "value": int(r.jobs), "salary": int(r.salary)} for r in titles.itertuples()],
        "levels": [{"label": LEVELS[k], "value": int(levels[k])} for k in ("EN", "MI", "SE", "EX") if k in levels.index],
        "skills": [{"label": k, "value": int(v), "share": round(v / len(d) * 100)} for k, v in skills.items()],
        "industries": [{"label": k, "value": int(v)} for k, v in industries.items()],
        "modes": [{"label": k, "value": int(v)} for k, v in modes.items()],
        "education": [{"label": k, "value": int(edu[k])} for k in EDU_RANK if k in edu.index],
        "months": [{"label": k, "value": int(v)} for k, v in months.items()],
    }
    if not country:
        top = (DF.groupby("company_location").agg(salary=("salary_usd", "mean"), jobs=("job_id", "count"))
               .sort_values("salary", ascending=False).head(8).reset_index())
        out["countries"] = [{"label": r.company_location, "value": int(r.salary), "jobs": int(r.jobs)}
                            for r in top.itertuples()]
    return jsonify(out)


# ----------------------------------------------------------------------------
# Role compare API: up to three roles side by side
# ----------------------------------------------------------------------------
@app.get("/api/compare")
def api_compare():
    roles = []
    for r in request.args.get("roles", "").split(","):
        if r in ALL_TITLES and r not in roles:
            roles.append(r)
    roles = roles[:3]
    country = request.args.get("country", "").strip()
    d = DF[DF["company_location"] == country] if country else DF
    out = []
    for role in roles:
        g = d[d["job_title"] == role]
        if g.empty:
            out.append({"role": role, "jobs": 0})
            continue
        lv = g.groupby("experience_level")["salary_usd"].mean()
        sk = g["skills_list"].explode().value_counts().head(6)
        out.append({
            "role": role, "jobs": int(len(g)),
            "avg_salary": int(g["salary_usd"].mean()), "median_salary": int(g["salary_usd"].median()),
            "avg_years": round(float(g["years_experience"].mean()), 1),
            "remote_pct": round(float((g["remote_ratio"] == 100).mean() * 100)),
            "benefits": round(float(g["benefits_score"].mean()), 1),
            "top_education": g["education_required"].value_counts().index[0],
            "top_industries": list(g["industry"].value_counts().head(3).index),
            "levels": [{"label": LEVELS[k], "value": int(lv[k])} for k in ("EN", "MI", "SE", "EX") if k in lv.index],
            "skills": [{"label": k, "share": round(v / len(g) * 100)} for k, v in sk.items()],
        })
    return jsonify({"country": country or "All countries", "roles": out})


# ----------------------------------------------------------------------------
# Resume reader
# ----------------------------------------------------------------------------
SKILL_PATTERNS = {
    "Python": r"python|pandas|numpy|scikit[- ]?learn|sklearn|jupyter|matplotlib",
    "SQL": r"sql|mysql|postgres(?:ql)?|sqlite|oracle",
    "TensorFlow": r"tensorflow|keras|tf\.keras",
    "PyTorch": r"pytorch|torch",
    "Kubernetes": r"kubernetes|k8s",
    "Scala": r"scala",
    "Linux": r"linux|ubuntu|bash|shell scripting",
    "Git": r"git|github|gitlab|bitbucket",
    "Java": r"java",
    "GCP": r"gcp|google cloud",
    "Hadoop": r"hadoop|hdfs|hive",
    "Tableau": r"tableau|power ?bi",
    "Computer Vision": r"computer vision|opencv|image processing|object detection|yolo|cnn",
    "Data Visualization": r"data visuali[sz]ation|visuali[sz]ation|seaborn|plotly|dashboards?",
    "Spark": r"spark|pyspark",
    "MLOps": r"mlops|ml ops|mlflow|kubeflow|model deployment|ci/cd",
    "Azure": r"azure",
    "Deep Learning": r"deep learning|neural networks?|lstm|rnn|transformers?|gans?|machine learning",
    "NLP": r"nlp|natural language processing|text mining|sentiment analysis|bert|llms?|chatbots?",
    "AWS": r"aws|amazon web services|ec2|s3|sagemaker",
    "Mathematics": r"mathematics|linear algebra|calculus|discrete math",
    "Docker": r"docker|containeri[sz]ation",
    "Statistics": r"statistics|statistical|probability|hypothesis testing|regression",
}
SKILL_REGEX = {
    name: re.compile(r"(?<![A-Za-z0-9+#])(?:" + pat + r")(?![A-Za-z0-9+#])", re.I)
    for name, pat in SKILL_PATTERNS.items()
}
# "R" is a one-letter language, so only count it in clear programming contexts
R_REGEX = re.compile(
    r"(?:\bR\s+(?:programming|language|studio)\b|\bRStudio\b|\bggplot2?\b|\btidyverse\b"
    r"|(?<=[,;|\u2022\u00b7/:(])\s*R\s*(?=[,;|\u2022\u00b7/)\n]|$))", re.M)

COUNTRY_HINTS = {
    "India": r"india|kerala|tamil nadu|karnataka|maharashtra|delhi|mumbai|bangalore|bengaluru|chennai|hyderabad|kochi|cochin|ernakulam|thiruvananthapuram|trivandrum|kozhikode|calicut|thrissur|kottayam|muvattupuzha|vazhakulam|pune|kolkata|\+91",
    "United States": r"united states|u\.s\.a?\b|usa|new york|san francisco|california|texas|seattle|boston|chicago",
    "United Kingdom": r"united kingdom|england|london|manchester|scotland|\+44",
    "Canada": r"canada|toronto|vancouver|montreal|ottawa",
    "Australia": r"australia|sydney|melbourne|brisbane",
    "Germany": r"germany|berlin|munich|hamburg|\+49",
    "France": r"france|paris",
    "Singapore": r"singapore",
    "Japan": r"japan|tokyo|osaka",
    "South Korea": r"south korea|seoul",
    "China": r"china|beijing|shanghai|shenzhen",
    "Ireland": r"ireland|dublin",
    "Netherlands": r"netherlands|amsterdam|rotterdam",
    "Sweden": r"sweden|stockholm",
    "Norway": r"norway|oslo",
    "Denmark": r"denmark|copenhagen",
    "Finland": r"finland|helsinki",
    "Switzerland": r"switzerland|zurich|geneva",
    "Austria": r"austria|vienna",
    "Israel": r"israel|tel aviv",
}
COUNTRY_REGEX = {c: re.compile(r"(?<![A-Za-z])(?:" + p + r")(?![A-Za-z])", re.I) for c, p in COUNTRY_HINTS.items()}

EDU_PATTERNS = [
    ("PhD", re.compile(r"\bph\.?\s?d\b|doctorate|doctor of philosophy", re.I), None),
    ("Master", re.compile(r"\bmaster'?s?\b|\bm\.?\s?tech\b|\bm\.?\s?sc\b|\bmca\b|\bmba\b", re.I),
     re.compile(r"\bM\.E\b|\bME\b")),
    ("Bachelor", re.compile(r"\bbachelor'?s?\b|\bb\.?\s?tech\b|\bb\.?\s?sc\b|\bbca\b|\bb\.?\s?com\b", re.I),
     re.compile(r"\bB\.E\b|\bBE\b")),
    ("Associate", re.compile(r"\bassociate degree\b|\bdiploma\b|\bassociate\b", re.I), None),
]
FIELDS = ["Artificial Intelligence", "Data Science", "Computer Science", "Information Technology",
          "Electronics", "Electrical", "Mechanical", "Civil", "Statistics", "Mathematics", "Business"]

SECTION_WORDS = re.compile(
    r"resume|curriculum|vitae|\bcv\b|profile|summary|objective|contact|education|experience|skills|projects", re.I)


def read_resume_text(filename, data):
    """Return plain text from a PDF, DOCX or TXT file."""
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Please upload a PDF, DOCX or TXT file.")


def find_name(lines, email):
    for ln in lines[:8]:
        ln = ln.strip()
        if not ln or "@" in ln or re.search(r"\d", ln) or SECTION_WORDS.search(ln):
            continue
        words = ln.split()
        if 1 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in words):
            return " ".join(w.capitalize() if w.isupper() and len(w) > 2 else w for w in words)
    if email:
        return re.sub(r"[\d_.]+", " ", email.split("@")[0]).strip().title()
    return ""


def find_country(text):
    """Pick the country mentioned earliest in the resume (the header usually has the address)."""
    best, best_pos = "", None
    for country, rx in COUNTRY_REGEX.items():
        m = rx.search(text)
        if m and (best_pos is None or m.start() < best_pos):
            best, best_pos = country, m.start()
    return best


def find_education(text):
    for level, rx, rx_case in EDU_PATTERNS:  # highest level first
        if rx.search(text) or (rx_case and rx_case.search(text)):
            return level
    return ""


def suggest_roles(text, skills):
    found = []
    low = text.lower()
    for title in ALL_TITLES:
        pos = low.find(title.lower())
        if pos >= 0:
            found.append((pos, title))
    roles = [t for _, t in sorted(found)]
    s = set(skills)
    guess = []
    if "Computer Vision" in s: guess.append("Computer Vision Engineer")
    if "NLP" in s: guess.append("NLP Engineer")
    if "MLOps" in s: guess.append("ML Ops Engineer")
    if s & {"Deep Learning", "TensorFlow", "PyTorch"}: guess.append("Machine Learning Engineer")
    if s & {"Spark", "Hadoop", "Scala"}: guess.append("Data Engineer")
    if s & {"Tableau", "Data Visualization", "Statistics"}: guess.append("Data Analyst")
    guess.append("Data Scientist")
    for g in guess:
        if g not in roles:
            roles.append(g)
    return roles[:3]


def extract_resume_info(text):
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    email = (re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text) or [None])[0] or ""
    phone = ""
    for m in re.finditer(r"\+?\d[\d\s\-().]{8,}\d", text):
        if len(re.sub(r"\D", "", m.group())) >= 10:
            phone = m.group().strip()
            break
    links = []
    for m in re.finditer(r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com)/[\w\-./]+", text, re.I):
        links.append(m.group().rstrip("."))

    skills = [name for name, rx in SKILL_REGEX.items() if rx.search(text)]
    if R_REGEX.search(text):
        skills.append("R")
    skills = sorted(set(skills))

    years = 0.0
    for m in re.finditer(r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)", text, re.I):
        ctx = text[max(0, m.start() - 60): m.end() + 60].lower()
        if "experience" in ctx and float(m.group(1)) <= 40:
            years = max(years, float(m.group(1)))

    institutions = []
    for ln in lines:
        if re.search(r"university|college|institute|school of", ln, re.I) and len(ln) < 120:
            if ln not in institutions:
                institutions.append(ln)
    field = next((f for f in FIELDS if re.search(re.escape(f), text, re.I)), "")

    roles = suggest_roles(text, skills)
    return {
        "name": find_name(lines, email), "email": email, "phone": phone, "links": links[:3],
        "country": find_country(text), "education": find_education(text), "field": field,
        "institutions": institutions[:3], "years": years, "skills": skills,
        "roles": roles, "role": roles[0] if roles else "", "words": len(text.split()),
    }


@app.post("/api/extract")
def api_extract():
    f = request.files.get("resume")
    if not f or not f.filename:
        return jsonify({"error": "Please choose a resume file first."}), 400
    try:
        text = read_resume_text(f.filename, f.read())
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception:
        return jsonify({"error": "Sorry, this file could not be read. Try another PDF or DOCX."}), 400
    if len(text.strip()) < 30:
        return jsonify({"error": "No readable text found. Scanned images are not supported; try a text-based PDF or DOCX."}), 400
    return jsonify(extract_resume_info(text))


@app.errorhandler(413)
def too_large(_):
    return jsonify({"error": "File is too large (limit 8 MB)."}), 413


if __name__ == "__main__":
    app.run(debug=True, port=5000)
